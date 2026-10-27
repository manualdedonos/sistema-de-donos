#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gera a planilha "Plano de Receita" (.xlsx) do Manual de Donos a partir de um JSON.

Uso:
    python3 gerar_plano_receita_xlsx.py dados.json "Plano_de_Receita_<Negocio>.xlsx"
    python3 gerar_plano_receita_xlsx.py --modelo "Plano_de_Receita_MODELO.xlsx"

Requer: Python 3 e openpyxl.

A planilha sai com fórmulas. Elas são calculadas quando o arquivo é aberto no
Excel, Google Planilhas ou LibreOffice (o arquivo vai marcado para recalcular
ao abrir). Não é preciso nenhum script de recálculo de outra skill. Para ter
os números já calculados (para o chat ou para conferir), o gerador imprime o
resumo do plano no final, usando calcular_plano.py (mesmas fórmulas, só
Python, sem openpyxl). Para calcular sem gerar a planilha:
    python3 calcular_plano.py dados.json

----------------------------------------------------------------------
FORMATO DO JSON DE ENTRADA
----------------------------------------------------------------------
Só "fontes" é obrigatório. Os demais campos são opcionais.

{
  "negocio": "Ativa Consultoria",
  "modelo_negocio": "Consultoria de gestao financeira para PMEs (servico B2B, projeto)",
  "periodo": "Mês de referência: julho/2026",
  "meta_receita": 450000,             # meta de receita para o MESMO período
  "crescimento_padrao": 0.03,         # crescimento mensal padrão da projeção (0.03 = 3%)
  "nome_entrada": "Entradas",         # como o negócio chama o contato inicial
  "nome_proposta": "Propostas",       # como o negócio chama o meio do funil
  "fontes": [                         # OBRIGATÓRIO — 1 linha por fonte de receita
    {
      "fonte": "Indicação de clientes",
      "entradas": 8,                  # quantos contatos iniciais chegaram
      "propostas": 6,                 # quantos viraram proposta/atendimento (0 se não mede)
      "vendas": 4,                    # quantas vendas fecharam
      "receita": 100000               # quanto de receita essa fonte gerou no período
    }
  ],
  "processos": [                      # opcional — atividades que fazem a venda acontecer
    {
      "fonte": "Indicação de clientes",
      "atividade": "Pedir indicação ao fim de todo projeto entregue",
      "padrao": "No call de encerramento, pedir 2 nomes e o contexto de cada um",
      "responsavel": "Marina Alves",
      "meta_atividade": "100% dos projetos encerrados",
      "frequencia": "A cada entrega",
      "verificacao": "Campo 'indicações pedidas' preenchido no CRM"
    }
  ]
}
----------------------------------------------------------------------
"""

import json
import math
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------
# Identidade visual — Manual de Donos
# --------------------------------------------------------------------------
LARANJA = "FF4500"
PRETO = "111111"
OFFWHITE = "F7F5F0"
CINZA = "E8E6DC"
CINZA_ESC = "8A8880"
AMARELO = "FFF2CC"      # células que o dono preenche
VERDE_CL = "E2EFDA"
AZUL_CL = "DDEBF7"

AZUL_INPUT = "0000FF"   # texto azul = input digitado
VERDE_LINK = "006400"   # texto verde = veio de outra aba

FONTE = "Arial"

BORDA_FINA = Side(style="thin", color="BFBFBF")
GRADE = Border(left=BORDA_FINA, right=BORDA_FINA, top=BORDA_FINA, bottom=BORDA_FINA)

FMT_RS = '"R$" #,##0'
FMT_RS0 = '"R$" #,##0;-"R$" #,##0;"-"'
FMT_RS2 = '"R$" #,##0.00'
FMT_PCT = '0.0%'
FMT_INT = '#,##0'

MAX_FONTES = 12          # linhas de fonte na aba 1
LINHA_INI = 9            # primeira linha de dados da aba 1
NOVOS_CANAIS = 3         # linhas extras de canal novo no simulador


# --------------------------------------------------------------------------
# Helpers de estilo
# --------------------------------------------------------------------------
def cel(ws, ref, valor, *, bold=False, size=10, color=PRETO, fill=None,
        fmt=None, align=None, wrap=False, italic=False, border=False,
        valign="center"):
    c = ws[ref]
    c.value = valor
    c.font = Font(name=FONTE, bold=bold, size=size, color=color, italic=italic)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(horizontal=align or "left", vertical=valign, wrap_text=wrap)
    if border:
        c.border = GRADE
    return c


def titulo_aba(ws, texto, subtitulo, ultima_col):
    ws.merge_cells(f"A1:{ultima_col}1")
    cel(ws, "A1", texto, bold=True, size=18, color="FFFFFF", fill=PRETO, align="left")
    ws.row_dimensions[1].height = 34
    ws.merge_cells(f"A2:{ultima_col}2")
    cel(ws, "A2", subtitulo, bold=True, size=10, color="FFFFFF", fill=LARANJA, align="left")
    ws.row_dimensions[2].height = 20


def faixa(ws, row, texto, ultima_col, cor=PRETO):
    ws.merge_cells(f"A{row}:{ultima_col}{row}")
    cel(ws, f"A{row}", texto, bold=True, size=11, color="FFFFFF", fill=cor, align="left")
    ws.row_dimensions[row].height = 22


def paragrafo(ws, row, texto, ultima_col, *, bold=False, italic=False,
              color=PRETO, fill=None, chars=None):
    ws.merge_cells(f"A{row}:{ultima_col}{row}")
    cel(ws, f"A{row}", texto, bold=bold, italic=italic, size=10, color=color,
        fill=fill, wrap=True, align="left", valign="top")
    largura = chars or 150
    linhas = max(1, math.ceil(len(texto) / largura))
    ws.row_dimensions[row].height = max(16, linhas * 14)


def cabecalho_tabela(ws, row, headers, col_ini=1):
    for i, h in enumerate(headers):
        ref = f"{get_column_letter(col_ini + i)}{row}"
        cel(ws, ref, h, bold=True, size=9, color="FFFFFF", fill=PRETO,
            align="center", wrap=True, border=True)
    ws.row_dimensions[row].height = 32


def larguras(ws, mapa):
    for col, w in mapa.items():
        ws.column_dimensions[col].width = w


def input_style(c):
    c.font = Font(name=FONTE, bold=True, size=10, color=AZUL_INPUT)
    c.fill = PatternFill("solid", fgColor=AMARELO)


def link_style(c):
    c.font = Font(name=FONTE, size=10, color=VERDE_LINK)


# --------------------------------------------------------------------------
# ABA 0 — Comece aqui
# --------------------------------------------------------------------------
def aba_comece_aqui(wb, dados):
    ws = wb.create_sheet("Comece aqui")
    ws.sheet_view.showGridLines = False
    larguras(ws, {"A": 26, "B": 34, "C": 34, "D": 34, "E": 22, "F": 18})

    titulo_aba(ws, "PLANO DE RECEITA", "MANUAL DE DONOS · de onde a venda vem, como ela acontece e onde colocar o próximo esforço", "F")

    r = 4
    paragrafo(ws, r, "Gestão de receita é conseguir olhar para a receita e saber de onde ela vem, "
                     "onde mexer e como gerar mais. Enquanto o faturamento for um bolo só, você "
                     "reage à receita. Quando cada fonte tem nome e número, você opera a receita.", "F", chars=120)
    r += 2

    ws.merge_cells(f"A{r}:F{r}")
    cel(ws, f"A{r}", "A PERGUNTA QUE GUIA TUDO:  O QUE precisa ser feito, QUANTAS vezes, COMO e POR QUEM, para vendermos?",
        bold=True, size=12, color=LARANJA, fill=OFFWHITE, align="center")
    ws.row_dimensions[r].height = 30
    r += 2

    faixa(ws, r, "OS 3 NÍVEIS DO PLANO DE RECEITA", "F")
    r += 1
    cabecalho_tabela(ws, r, ["Nível", "Pergunta-chave", "O que você passa a saber", "O que isso destrava", "Onde está na planilha"])
    r += 1
    niveis = [
        ["NÍVEL 1\nVisibilidade", "De onde a venda vem?",
         "Quanto cada fonte de receita gera.",
         "Você para de olhar o faturamento como um bolo só e enxerga dinheiros diferentes.",
         "Aba 1 · colunas Receita e % da receita"],
        ["NÍVEL 2\nControle", "Como cada fonte vira venda?",
         "Quantas entradas viram proposta e quantas propostas viram venda, por fonte.",
         "Você enxerga o funil de cada fonte e acha o gargalo — o ponto exato onde a venda trava.",
         "Aba 1 · colunas Entradas, Propostas, Vendas e Conversões"],
        ["NÍVEL 3\nDecisão", "Onde colocar o próximo esforço?",
         "Quanto de receita cada entrada gera, fonte por fonte.",
         "Você descobre que R$ 1 de um canal não vale o mesmo que R$ 1 de outro — e para de investir tempo no lugar errado.",
         "Aba 1 · coluna Receita por entrada"],
    ]
    for n in niveis:
        for i, v in enumerate(n):
            cel(ws, f"{get_column_letter(1+i)}{r}", v, size=9, wrap=True, border=True,
                bold=(i == 0), color=(LARANJA if i == 0 else PRETO), valign="top")
        ws.row_dimensions[r].height = 52
        r += 1
    r += 1

    faixa(ws, r, "COMO USAR ESSA PLANILHA — 6 PASSOS", "F")
    r += 1
    passos = [
        ("1. Nomeie suas fontes", "Na aba «1. Plano de Receita», escreva todas as fontes de receita que JÁ existem hoje. Não invente nenhuma nova ainda."),
        ("2. Preencha o último mês fechado", "Para cada fonte: entradas, propostas, vendas e receita. Se você não mede alguma coisa, escreva sua melhor estimativa e marque para começar a contar. Aprender a contar vem antes de otimizar."),
        ("3. Leia o diagnóstico", "A própria aba 1 aponta sua maior fonte, seu melhor canal por entrada, seu maior gargalo de conversão e onde você gasta esforço sem retorno."),
        ("4. Simule a meta", "Na aba «2. Simulador de Meta», coloque a meta de receita e mexa nas entradas, conversões e ticket. A planilha mostra o gap e quantas entradas faltam no melhor canal para fechá-lo."),
        ("5. Projete o ano", "A aba «3. Projeção 12 meses» transforma o mês simulado em um plano de 12 meses, fonte por fonte."),
        ("6. Transforme em processo e cadência", "Na aba «5. Processos e Cadência», escreva o que precisa ser feito, quantas vezes, como e por quem — e marque o comitê de receita semanal na agenda. Plano sem cadência vira intenção."),
    ]
    for tit, txt in passos:
        cel(ws, f"A{r}", tit, bold=True, size=10, color=LARANJA, wrap=True, valign="top")
        ws.merge_cells(f"B{r}:F{r}")
        cel(ws, f"B{r}", txt, size=10, wrap=True, valign="top")
        ws.row_dimensions[r].height = max(16, math.ceil(len(txt) / 105) * 14)
        r += 1
    r += 1

    faixa(ws, r, "LEGENDA — o que você edita e o que a planilha calcula", "F", cor=LARANJA)
    r += 1
    legendas = [
        ("Texto azul, fundo amarelo", "Você preenche. É input seu.", AMARELO, AZUL_INPUT),
        ("Texto preto", "Fórmula. Não mexa — a planilha calcula sozinha.", None, PRETO),
        ("Texto verde", "Veio de outra aba. Se mudar lá, muda aqui.", None, VERDE_LINK),
    ]
    for amostra, expl, fill, color in legendas:
        c = cel(ws, f"A{r}", amostra, bold=(color == AZUL_INPUT), size=10, color=color, border=True)
        if fill:
            c.fill = PatternFill("solid", fgColor=fill)
        ws.merge_cells(f"B{r}:F{r}")
        cel(ws, f"B{r}", expl, size=10)
        r += 1
    r += 1

    paragrafo(ws, r, "REGRA DE OURO: você não gerencia o resultado, você gerencia as atividades e monitora "
                     "os indicadores de direção — o resultado é consequência. Crescer é empilhar fontes que "
                     "funcionam: acertou uma, mantenha; vá buscar a próxima; some forças, não troque apostas.",
              "F", bold=True, fill=OFFWHITE, chars=120)
    r += 2
    paragrafo(ws, r, "Baseado no manual «Processo comercial: manual para destravar as vendas» — Manual de Donos, com Jonathan Souza.",
              "F", italic=True, color=CINZA_ESC, chars=150)
    return ws


# --------------------------------------------------------------------------
# ABA 1 — Plano de Receita
# --------------------------------------------------------------------------
def aba_plano(wb, dados):
    ws = wb.create_sheet("1. Plano de Receita")
    ws.sheet_view.showGridLines = False
    larguras(ws, {"A": 30, "B": 12, "C": 12, "D": 11, "E": 15, "F": 13,
                  "G": 13, "H": 13, "I": 13, "J": 15, "K": 12})

    nome_entrada = dados.get("nome_entrada", "Entradas")
    nome_proposta = dados.get("nome_proposta", "Propostas")

    titulo_aba(ws, "1. PLANO DE RECEITA — O RETRATO DE HOJE",
               "Níveis 1, 2 e 3 na mesma tabela: de onde vem, como acontece e quanto vale cada entrada", "K")

    cel(ws, "A4", "Negócio:", bold=True)
    input_style(cel(ws, "B4", dados.get("negocio", "")))
    ws.merge_cells("B4:E4")
    cel(ws, "G4", "Período de referência:", bold=True)
    input_style(cel(ws, "I4", dados.get("periodo", "Último mês fechado")))
    ws.merge_cells("I4:K4")

    cel(ws, "A5", "Modelo de negócio:", bold=True)
    input_style(cel(ws, "B5", dados.get("modelo_negocio", "")))
    ws.merge_cells("B5:K5")

    faixa(ws, 7, "PREENCHA AS COLUNAS AMARELAS. O RESTO A PLANILHA CALCULA.", "K", cor=LARANJA)

    headers = ["Fonte de receita", nome_entrada, nome_proposta, "Vendas", "Receita (R$)",
               "Ticket médio", "Conv. entrada→proposta", "Conv. proposta→venda",
               "Conv. Entrada→Venda", "Receita por entrada", "% da receita"]
    cabecalho_tabela(ws, 8, headers)

    fontes = dados.get("fontes") or []
    fim = LINHA_INI + MAX_FONTES - 1

    for i in range(MAX_FONTES):
        r = LINHA_INI + i
        f = fontes[i] if i < len(fontes) else None
        exemplo = (f is None and i == 0 and not fontes)

        if f:
            nome, ent, prop, ven, rec = (f.get("fonte", ""), f.get("entradas"),
                                         f.get("propostas"), f.get("vendas"), f.get("receita"))
        elif exemplo:
            nome, ent, prop, ven, rec = ("Indicação  (exemplo — apague e coloque as suas)", 40, 20, 11, 80000)
        else:
            nome, ent, prop, ven, rec = ("", None, None, None, None)

        c = cel(ws, f"A{r}", nome, border=True)
        input_style(c)
        if exemplo:
            c.font = Font(name=FONTE, italic=True, size=10, color=CINZA_ESC)
        for col, val, fmt in (("B", ent, FMT_INT), ("C", prop, FMT_INT),
                              ("D", ven, FMT_INT), ("E", rec, FMT_RS)):
            cc = cel(ws, f"{col}{r}", val, fmt=fmt, align="center", border=True)
            input_style(cc)
            cc.number_format = fmt
            if exemplo:
                cc.font = Font(name=FONTE, italic=True, size=10, color=CINZA_ESC)

        cel(ws, f"F{r}", f"=IFERROR(E{r}/D{r},\"\")", fmt=FMT_RS, align="center", border=True)
        cel(ws, f"G{r}", f"=IFERROR(C{r}/B{r},\"\")", fmt=FMT_PCT, align="center", border=True)
        cel(ws, f"H{r}", f"=IFERROR(D{r}/C{r},\"\")", fmt=FMT_PCT, align="center", border=True)
        cel(ws, f"I{r}", f"=IFERROR(D{r}/B{r},\"\")", fmt=FMT_PCT, align="center", border=True)
        cel(ws, f"J{r}", f"=IFERROR(E{r}/B{r},\"\")", fmt=FMT_RS2, align="center", border=True)
        cel(ws, f"K{r}", f"=IFERROR(E{r}/$E${fim+1},\"\")", fmt=FMT_PCT, align="center", border=True)

    tot = fim + 1
    cel(ws, f"A{tot}", "TOTAL", bold=True, color="FFFFFF", fill=PRETO, border=True)
    for col in ("B", "C", "D", "E"):
        cel(ws, f"{col}{tot}", f"=SUM({col}{LINHA_INI}:{col}{fim})", bold=True, color="FFFFFF",
            fill=PRETO, align="center", border=True,
            fmt=(FMT_RS if col == "E" else FMT_INT))
    cel(ws, f"F{tot}", f"=IFERROR(E{tot}/D{tot},\"\")", bold=True, color="FFFFFF", fill=PRETO,
        align="center", border=True, fmt=FMT_RS)
    cel(ws, f"G{tot}", f"=IFERROR(C{tot}/B{tot},\"\")", bold=True, color="FFFFFF", fill=PRETO,
        align="center", border=True, fmt=FMT_PCT)
    cel(ws, f"H{tot}", f"=IFERROR(D{tot}/C{tot},\"\")", bold=True, color="FFFFFF", fill=PRETO,
        align="center", border=True, fmt=FMT_PCT)
    cel(ws, f"I{tot}", f"=IFERROR(D{tot}/B{tot},\"\")", bold=True, color="FFFFFF", fill=PRETO,
        align="center", border=True, fmt=FMT_PCT)
    cel(ws, f"J{tot}", f"=IFERROR(E{tot}/B{tot},\"\")", bold=True, color="FFFFFF", fill=PRETO,
        align="center", border=True, fmt=FMT_RS2)
    cel(ws, f"K{tot}", f"=IFERROR(SUM(K{LINHA_INI}:K{fim}),\"\")", bold=True,
        color="FFFFFF", fill=PRETO, align="center", border=True, fmt=FMT_PCT)

    r = tot + 2
    faixa(ws, r, "LEITURA AUTOMÁTICA DO SEU PLANO — o que esses números já estão te dizendo", "K")
    r += 1
    leituras = [
        ("Sua maior fonte de receita",
         f"=IFERROR(INDEX($A${LINHA_INI}:$A${fim},MATCH(MAX($E${LINHA_INI}:$E${fim}),$E${LINHA_INI}:$E${fim},0)),\"\")",
         f"=IFERROR(MAX($E${LINHA_INI}:$E${fim}),\"\")", FMT_RS,
         "É de onde o dinheiro vem hoje. Antes de buscar canal novo, garanta que essa fonte não vai cair."),
        ("Sua entrada mais valiosa (maior receita por entrada)",
         f"=IFERROR(INDEX($A${LINHA_INI}:$A${fim},MATCH(MAX($J${LINHA_INI}:$J${fim}),$J${LINHA_INI}:$J${fim},0)),\"\")",
         f"=IFERROR(MAX($J${LINHA_INI}:$J${fim}),\"\")", FMT_RS2,
         "Se você tem 1 hora para vender, é aqui que ela rende mais. Pergunte-se: como eu gero MAIS entradas nesse canal?"),
        ("Sua entrada menos valiosa (menor receita por entrada)",
         f"=IFERROR(INDEX($A${LINHA_INI}:$A${fim},MATCH(MIN($J${LINHA_INI}:$J${fim}),$J${LINHA_INI}:$J${fim},0)),\"\")",
         f"=IFERROR(MIN($J${LINHA_INI}:$J${fim}),\"\")", FMT_RS2,
         "Não significa abandonar. Significa que aqui o jogo é volume e eficiência — e que uma melhora pequena por entrada vira muito dinheiro."),
        ("Onde você gasta mais esforço (mais entradas)",
         f"=IFERROR(INDEX($A${LINHA_INI}:$A${fim},MATCH(MAX($B${LINHA_INI}:$B${fim}),$B${LINHA_INI}:$B${fim},0)),\"\")",
         f"=IFERROR(MAX($B${LINHA_INI}:$B${fim}),\"\")", FMT_INT,
         "Compare com a linha acima: se o maior esforço estiver na entrada menos valiosa, seu time está trabalhando no lugar errado."),
        ("Seu maior gargalo (pior conversão entrada→venda)",
         f"=IFERROR(INDEX($A${LINHA_INI}:$A${fim},MATCH(MIN($I${LINHA_INI}:$I${fim}),$I${LINHA_INI}:$I${fim},0)),\"\")",
         f"=IFERROR(MIN($I${LINHA_INI}:$I${fim}),\"\")", FMT_PCT,
         "É o ponto do funil onde a venda trava. Arrumar conversão é mais barato do que comprar mais entradas."),
    ]
    for rotulo, f_nome, f_valor, fmt, expl in leituras:
        cel(ws, f"A{r}", rotulo, bold=True, size=9, wrap=True, border=True, valign="top")
        ws.merge_cells(f"B{r}:D{r}")
        cel(ws, f"B{r}", f_nome, bold=True, size=10, color=LARANJA, border=True, wrap=True)
        cel(ws, f"E{r}", f_valor, bold=True, size=10, fmt=fmt, align="center", border=True)
        ws.merge_cells(f"F{r}:K{r}")
        cel(ws, f"F{r}", expl, size=9, wrap=True, italic=True, valign="top", border=True)
        ws.row_dimensions[r].height = 30
        r += 1

    r += 1
    faixa(ws, r, "COMO PREENCHER CADA COLUNA", "K", cor=LARANJA)
    r += 1
    ajuda = [
        ("Fonte de receita", "Dê nome a cada origem de venda: carteira ativa, indicação, prospecção ativa, inbound, eventos, parcerias, canal indireto, fluxo de loja, WhatsApp, marketplace. Cada fonte é, na prática, um dinheiro diferente."),
        (nome_entrada, "Todo contato inicial do período: pessoa que entrou na loja, lead que preencheu formulário, orçamento pedido, visita agendada, DM respondida. Se você não sabe, comece contando — engenheiro de tráfego conta carro por carro antes de projetar a rua."),
        (nome_proposta, "O meio do caminho: proposta enviada, orçamento passado, atendimento com prova, cliente que pediu preço. Se o seu negócio não tem essa etapa, repita o número de vendas."),
        ("Vendas", "Quantos negócios fecharam no período, vindos daquela fonte."),
        ("Receita (R$)", "Quanto de receita aquela fonte gerou no período. A soma tem que bater com o seu faturamento total."),
    ]
    for rot, txt in ajuda:
        cel(ws, f"A{r}", rot, bold=True, size=9, wrap=True, valign="top")
        ws.merge_cells(f"B{r}:K{r}")
        cel(ws, f"B{r}", txt, size=9, wrap=True, valign="top")
        ws.row_dimensions[r].height = max(16, math.ceil(len(txt) / 130) * 13)
        r += 1
    return ws


# --------------------------------------------------------------------------
# ABA 2 — Simulador de Meta
# --------------------------------------------------------------------------
def aba_simulador(wb, dados):
    ws = wb.create_sheet("2. Simulador de Meta")
    ws.sheet_view.showGridLines = False
    larguras(ws, {"A": 30, "B": 15, "C": 15, "D": 14, "E": 13, "F": 16,
                  "G": 15, "H": 16, "I": 15, "J": 12})

    P = "'1. Plano de Receita'"
    fim1 = LINHA_INI + MAX_FONTES - 1

    titulo_aba(ws, "2. SIMULADOR DE META",
               "Mexa nas entradas, na conversão e no ticket. Veja a receita mudar e o gap até a meta fechar.", "J")

    cel(ws, "A4", "META DE RECEITA DO PERÍODO:", bold=True, size=11)
    ws.merge_cells("A4:B4")
    meta = dados.get("meta_receita")
    input_style(cel(ws, "C4", meta, fmt=FMT_RS, align="center", border=True))
    ws["C4"].number_format = FMT_RS
    cel(ws, "D4", "Receita de hoje (aba 1):", bold=True, align="right")
    ws.merge_cells("D4:E4")
    c = cel(ws, "F4", f"={P}!E{fim1+1}", fmt=FMT_RS, align="center", border=True, bold=True)
    link_style(c)
    c.font = Font(name=FONTE, bold=True, size=10, color=VERDE_LINK)

    faixa(ws, 6, "AS LINHAS COMEÇAM IGUAIS AO SEU RETRATO DE HOJE. SOBRESCREVA À VONTADE — É PARA ISSO QUE ELA SERVE.", "J", cor=LARANJA)

    headers = ["Fonte de receita", "Entradas planejadas", "Conv. entrada→venda",
               "Ticket médio (R$)", "Vendas projetadas", "Receita projetada (R$)",
               "Receita por entrada", "Receita hoje (R$)", "Variação (R$)", "Variação %"]
    cabecalho_tabela(ws, 7, headers)

    ini = 8
    fim = ini + MAX_FONTES + NOVOS_CANAIS - 1
    for i in range(MAX_FONTES + NOVOS_CANAIS):
        r = ini + i
        if i < MAX_FONTES:
            r1 = LINHA_INI + i
            c = cel(ws, f"A{r}", f'=IF({P}!A{r1}="","",{P}!A{r1})', border=True)
            link_style(c)
            for col, src, fmt in (("B", f'=IF({P}!B{r1}="","",{P}!B{r1})', FMT_INT),
                                  ("C", f'=IF({P}!I{r1}="","",{P}!I{r1})', FMT_PCT),
                                  ("D", f'=IF({P}!F{r1}="","",{P}!F{r1})', FMT_RS)):
                cc = cel(ws, f"{col}{r}", src, fmt=fmt, align="center", border=True)
                cc.fill = PatternFill("solid", fgColor=AMARELO)
                cc.font = Font(name=FONTE, bold=True, size=10, color=VERDE_LINK)
            cel(ws, f"H{r}", f'=IF({P}!E{r1}="","",{P}!E{r1})', fmt=FMT_RS0, align="center", border=True)
            ws[f"H{r}"].font = Font(name=FONTE, size=10, color=VERDE_LINK)
        else:
            n = i - MAX_FONTES + 1
            cc = cel(ws, f"A{r}", f"Canal novo {n}  (escreva aqui)", border=True)
            input_style(cc)
            cc.font = Font(name=FONTE, italic=True, size=10, color=CINZA_ESC)
            for col, fmt in (("B", FMT_INT), ("C", FMT_PCT), ("D", FMT_RS)):
                cc = cel(ws, f"{col}{r}", None, fmt=fmt, align="center", border=True)
                input_style(cc)
                cc.number_format = fmt
            cel(ws, f"H{r}", 0, fmt=FMT_RS, align="center", border=True)

        cel(ws, f"E{r}", f"=IFERROR(B{r}*C{r},\"\")", fmt='#,##0.0', align="center", border=True)
        cel(ws, f"F{r}", f"=IFERROR(E{r}*D{r},0)", fmt=FMT_RS0, align="center", border=True)
        cel(ws, f"G{r}", f"=IFERROR(F{r}/B{r},\"\")", fmt=FMT_RS2, align="center", border=True)
        cel(ws, f"I{r}", f"=IFERROR(F{r}-H{r},\"\")", fmt=FMT_RS0, align="center", border=True)
        cel(ws, f"J{r}", f"=IFERROR(I{r}/H{r},\"\")", fmt=FMT_PCT, align="center", border=True)

    tot = fim + 1
    cel(ws, f"A{tot}", "TOTAL PROJETADO", bold=True, color="FFFFFF", fill=PRETO, border=True)
    for col, fmt in (("B", FMT_INT), ("E", '#,##0.0'), ("F", FMT_RS), ("H", FMT_RS), ("I", FMT_RS)):
        cel(ws, f"{col}{tot}", f"=SUM({col}{ini}:{col}{fim})", bold=True, color="FFFFFF",
            fill=PRETO, align="center", border=True, fmt=fmt)
    for col in ("C", "D", "G", "J"):
        cel(ws, f"{col}{tot}", "", fill=PRETO, border=True)
    cel(ws, f"G{tot}", f"=IFERROR(F{tot}/B{tot},\"\")", bold=True, color="FFFFFF", fill=PRETO,
        align="center", border=True, fmt=FMT_RS2)
    cel(ws, f"J{tot}", f"=IFERROR(I{tot}/H{tot},\"\")", bold=True, color="FFFFFF", fill=PRETO,
        align="center", border=True, fmt=FMT_PCT)

    r = tot + 2
    faixa(ws, r, "RESULTADO DA SIMULAÇÃO", "J", cor=LARANJA)
    r += 1
    base = r
    linhas = [
        ("Receita projetada", f"=F{tot}", FMT_RS, "Soma da simulação acima."),
        ("Meta de receita", "=C4", FMT_RS, "O número que você escreveu lá em cima."),
        ("Gap até a meta", f"=C4-F{tot}", FMT_RS, "Positivo = ainda falta. Negativo ou zero = a simulação já bate a meta."),
        ("Situação", f'=IF(C4="","Coloque a meta lá em cima",IF(C4-F{tot}<=0,"META BATIDA nessa simulação","AINDA FALTA — ajuste entradas, conversão ou ticket"))', None, "Leitura rápida."),
        ("Melhor canal por receita/entrada nessa simulação",
         f"=IFERROR(INDEX($A${ini}:$A${fim},MATCH(MAX($G${ini}:$G${fim}),$G${ini}:$G${fim},0)),\"\")",
         None, "É onde cada entrada nova rende mais."),
        ("Receita por entrada desse canal", f"=IFERROR(MAX($G${ini}:$G${fim}),\"\")", FMT_RS2,
         "Quanto entra de receita a cada contato inicial gerado nele."),
        ("Entradas adicionais nesse canal para fechar o gap",
         f"=IF(C4=\"\",\"\",IF(C4-F{tot}<=0,0,IFERROR(ROUNDUP((C4-F{tot})/MAX($G${ini}:$G${fim}),0),\"\")))",
         FMT_INT, "Essa é a decisão que o plano gera: quantas entradas a mais, no canal certo, fecham a conta."),
    ]
    for rot, form, fmt, expl in linhas:
        cel(ws, f"A{r}", rot, bold=True, size=10, wrap=True, border=True, valign="top")
        ws.merge_cells(f"B{r}:C{r}")
        cel(ws, f"B{r}", form, bold=True, size=11, color=LARANJA, fmt=fmt, align="center", border=True, wrap=True)
        ws.merge_cells(f"D{r}:J{r}")
        cel(ws, f"D{r}", expl, size=9, italic=True, wrap=True, valign="top", border=True)
        ws.row_dimensions[r].height = 24
        r += 1

    r += 1
    faixa(ws, r, "AS 3 ALAVANCAS QUE VOCÊ TEM — e o custo de cada uma", "J")
    r += 1
    alavancas = [
        ("1. Mais entradas (volume)", "Gerar mais contatos iniciais no canal. É a alavanca mais óbvia e normalmente a mais cara: exige mais anúncio, mais prospecção, mais gente. Só puxe essa depois de olhar as outras duas."),
        ("2. Melhor conversão (funil)", "Transformar em venda o que já chegou. É a alavanca mais barata: a entrada já foi paga. Uma conversão de 16% para 18% no seu maior canal costuma valer mais do que um mês inteiro de anúncio novo."),
        ("3. Ticket maior (valor)", "Vender mais por venda: combo, item adicional, escopo maior, plano anual. Não custa entrada nenhuma — custa processo de oferta. «Você conhece nosso produto mais vendido?» é ticket, não é sorte."),
    ]
    for rot, txt in alavancas:
        cel(ws, f"A{r}", rot, bold=True, size=10, color=LARANJA, wrap=True, valign="top", border=True)
        ws.merge_cells(f"B{r}:J{r}")
        cel(ws, f"B{r}", txt, size=9, wrap=True, valign="top", border=True)
        ws.row_dimensions[r].height = max(18, math.ceil(len(txt) / 120) * 13)
        r += 1
    return ws


# --------------------------------------------------------------------------
# ABA 3 — Projeção 12 meses
# --------------------------------------------------------------------------
def aba_projecao(wb, dados):
    ws = wb.create_sheet("3. Projeção 12 meses")
    ws.sheet_view.showGridLines = False
    larguras(ws, {"A": 30, "B": 14})
    for i in range(12):
        ws.column_dimensions[get_column_letter(3 + i)].width = 12
    ws.column_dimensions["O"].width = 16

    S = "'2. Simulador de Meta'"
    sim_ini, sim_fim = 8, 8 + MAX_FONTES + NOVOS_CANAIS - 1

    titulo_aba(ws, "3. PROJEÇÃO 12 MESES",
               "O mês simulado virando plano de ano, fonte por fonte. Mês 1 = a simulação da aba 2.", "O")

    cel(ws, "A4", "Crescimento padrão por mês:", bold=True)
    ws.merge_cells("A4:B4")
    input_style(cel(ws, "C4", dados.get("crescimento_padrao", 0.0), fmt=FMT_PCT, align="center", border=True))
    ws["C4"].number_format = FMT_PCT
    ws.merge_cells("D4:O4")
    cel(ws, "D4", "Aplicado a todas as fontes. Você pode sobrescrever fonte a fonte na coluna B — nem todo canal cresce no mesmo ritmo.",
        size=9, italic=True)

    headers = ["Fonte de receita", "Cresc./mês"] + [f"Mês {i+1}" for i in range(12)] + ["Total 12 meses"]
    cabecalho_tabela(ws, 6, headers)

    ini = 7
    fim = ini + MAX_FONTES + NOVOS_CANAIS - 1
    for i in range(MAX_FONTES + NOVOS_CANAIS):
        r = ini + i
        rs = sim_ini + i
        c = cel(ws, f"A{r}", f'=IF({S}!A{rs}="","",{S}!A{rs})', border=True)
        link_style(c)
        cc = cel(ws, f"B{r}", "=$C$4", fmt=FMT_PCT, align="center", border=True)
        cc.fill = PatternFill("solid", fgColor=AMARELO)
        cc.font = Font(name=FONTE, bold=True, size=10, color=AZUL_INPUT)
        c = cel(ws, f"C{r}", f"={S}!F{rs}", fmt=FMT_RS0, align="center", border=True)
        link_style(c)
        for j in range(1, 12):
            col = get_column_letter(3 + j)
            ant = get_column_letter(2 + j)
            cel(ws, f"{col}{r}", f"={ant}{r}*(1+$B{r})", fmt=FMT_RS0, align="center", border=True)
        cel(ws, f"O{r}", f"=SUM(C{r}:N{r})", fmt=FMT_RS0, align="center", border=True, bold=True)

    tot = fim + 1
    cel(ws, f"A{tot}", "TOTAL", bold=True, color="FFFFFF", fill=PRETO, border=True)
    cel(ws, f"B{tot}", "", fill=PRETO, border=True)
    for j in range(13):
        col = get_column_letter(3 + j)
        cel(ws, f"{col}{tot}", f"=SUM({col}{ini}:{col}{fim})", bold=True, color="FFFFFF",
            fill=PRETO, align="center", border=True, fmt=FMT_RS)

    r = tot + 2
    faixa(ws, r, "COMO LER ESSA PROJEÇÃO", "O", cor=LARANJA)
    r += 1
    for txt in [
        "Renomeie «Mês 1, Mês 2...» para os meses do seu ciclo (Jan, Fev, Mar...). Mês 1 é o mês que você simulou na aba 2.",
        "Crescimento não é fé: é consequência de mais entradas, melhor conversão ou ticket maior. Se você colocou 5% ao mês numa fonte, escreva na aba 5 QUAL atividade vai produzir esses 5%.",
        "Fontes diferentes crescem em ritmos diferentes. Indicação cresce com processo; inbound cresce com investimento; carteira ativa cresce com cadência de contato. Ajuste a coluna B fonte a fonte.",
        "Se a projeção do ano parece impossível, ela está te dizendo algo útil: ou a meta precisa mudar, ou o plano precisa de uma fonte nova — não de mais esforço na mesma.",
    ]:
        paragrafo(ws, r, "•  " + txt, "O", chars=170)
        r += 1
    return ws


# --------------------------------------------------------------------------
# ABA 4 — Acompanhamento
# --------------------------------------------------------------------------
def aba_acompanhamento(wb, dados):
    ws = wb.create_sheet("4. Acompanhamento")
    ws.sheet_view.showGridLines = False
    larguras(ws, {"A": 30})
    for i in range(12):
        ws.column_dimensions[get_column_letter(2 + i)].width = 12

    PJ = "'3. Projeção 12 meses'"
    proj_ini = 7
    n = MAX_FONTES + NOVOS_CANAIS

    titulo_aba(ws, "4. ACOMPANHAMENTO — META VS. REALIZADO",
               "O material do comitê de receita semanal/mensal: onde o plano está acontecendo e onde não está", "M")

    def bloco(row, titulo, cor, gerar_valor, fmt, editavel=False, link=False):
        faixa(ws, row, titulo, "M", cor=cor)
        cabecalho_tabela(ws, row + 1, ["Fonte de receita"] + [f"Mês {i+1}" for i in range(12)])
        for i in range(n):
            r = row + 2 + i
            c = cel(ws, f"A{r}", f'=IF({PJ}!A{proj_ini+i}="","",{PJ}!A{proj_ini+i})', border=True)
            link_style(c)
            for j in range(12):
                col = get_column_letter(2 + j)
                cc = cel(ws, f"{col}{r}", gerar_valor(i, j, r), fmt=fmt, align="center", border=True)
                if editavel:
                    cc.fill = PatternFill("solid", fgColor=AMARELO)
                    cc.font = Font(name=FONTE, bold=True, size=10, color=AZUL_INPUT)
                elif link:
                    cc.font = Font(name=FONTE, size=10, color=VERDE_LINK)
        tr = row + 2 + n
        cel(ws, f"A{tr}", "TOTAL", bold=True, color="FFFFFF", fill=PRETO, border=True)
        for j in range(12):
            col = get_column_letter(2 + j)
            cel(ws, f"{col}{tr}", f"=SUM({col}{row+2}:{col}{row+1+n})", bold=True, color="FFFFFF",
                fill=PRETO, align="center", border=True, fmt=fmt)
        return tr

    fim_meta = bloco(4, "META POR FONTE (vem da projeção — aba 3)", PRETO,
                     lambda i, j, r: f"={PJ}!{get_column_letter(3+j)}{proj_ini+i}",
                     FMT_RS0, link=True)

    ini_real = fim_meta + 2
    fim_real = bloco(ini_real, "REALIZADO POR FONTE (você preenche mês a mês)", LARANJA,
                     lambda i, j, r: None, FMT_RS0, editavel=True)

    ini_pct = fim_real + 2
    fim_pct = bloco(ini_pct, "% DA META ATINGIDA (calculado)", PRETO,
                    lambda i, j, r: (lambda c, rr, rm: f'=IF({c}{rr}="","",IFERROR({c}{rr}/{c}{rm},""))')(
                        get_column_letter(2+j), r-(ini_pct-ini_real), r-(ini_pct-4)),
                    FMT_PCT)

    faixa_pct = f"B{ini_pct+2}:M{ini_pct+1+n}"
    ws.conditional_formatting.add(faixa_pct, CellIsRule(
        operator="greaterThanOrEqual", formula=["1"],
        fill=PatternFill("solid", fgColor=VERDE_CL)))
    ws.conditional_formatting.add(faixa_pct, CellIsRule(
        operator="between", formula=["0.9", "0.9999"],
        fill=PatternFill("solid", fgColor="FFF2CC")))
    ws.conditional_formatting.add(faixa_pct, CellIsRule(
        operator="lessThan", formula=["0.9"],
        fill=PatternFill("solid", fgColor="F8CBAD")))

    r = fim_pct + 2
    faixa(ws, r, "O QUE FAZER COM ESSA ABA", "M", cor=LARANJA)
    r += 1
    for txt in [
        "Verde = bateu. Amarelo = chegou perto (90% a 99%). Vermelho = abaixo de 90% e precisa de decisão, não de torcida.",
        "Essa é a tabela do comitê de receita semanal: 30 minutos, meta contra realizado por fonte, indicadores de direção e UMA decisão concreta por semana — o que muda, quem faz, até quando.",
        "Se a atividade aconteceu e o número não se mexeu, o problema não é esforço: é o processo. Reveja o padrão daquela fonte na aba 5.",
        "Cada semana é uma volta do ciclo definir → executar → medir → ajustar. O acumulado dessas voltas é o que gera resultado consistente.",
    ]:
        paragrafo(ws, r, "•  " + txt, "M", chars=170)
        r += 1
    return ws


# --------------------------------------------------------------------------
# ABA 5 — Processos e Cadência
# --------------------------------------------------------------------------
def aba_processos(wb, dados):
    ws = wb.create_sheet("5. Processos e Cadência")
    ws.sheet_view.showGridLines = False
    larguras(ws, {"A": 24, "B": 30, "C": 38, "D": 20, "E": 20, "F": 16, "G": 28})

    titulo_aba(ws, "5. PROCESSOS E CADÊNCIA",
               "O plano só vira receita quando vira atividade com padrão, responsável e verificação", "G")

    paragrafo(ws, 4, "Atividade é o que está no seu controle fazer. Fazer uma pessoa entrar na loja não está no seu controle; "
                     "atendê-la depois que entrou, sim. Uma atividade só vira processo de verdade quando tem 3 coisas: "
                     "um PADRÃO que qualquer pessoa consegue seguir desde o primeiro dia, um RESPONSÁVEL com nome e sobrenome, "
                     "e uma FORMA DE VERIFICAR se foi feito.", "G", chars=140)

    paragrafo(ws, 5, "«Vou fazer 5 ligações por dia» é controle.  «Vou vender mais» é só visibilidade.",
              "G", bold=True, color=LARANJA, chars=140)

    faixa(ws, 7, "O QUE PRECISA SER FEITO, QUANTAS VEZES, COMO E POR QUEM — comece pelas suas 2 ou 3 principais fontes", "G")
    cabecalho_tabela(ws, 8, ["Fonte de receita", "Atividade-chave", "Padrão (como exatamente é feito)",
                             "Responsável (nome e sobrenome)", "Meta de atividade (quantas vezes)",
                             "Frequência", "Como verificar se aconteceu"])

    exemplos = [
        ["Indicação  (exemplo)", "Pedir indicação ao fim de cada entrega",
         "No call de encerramento, pedir 2 nomes e o contexto de cada um antes de encerrar a chamada",
         "Marina Alves", "100% dos projetos encerrados", "A cada entrega",
         "Campo «indicações pedidas» preenchido no CRM"],
        ["Inbound  (exemplo)", "Primeiro atendimento do lead",
         "Responder todo lead em até 5 minutos, por telefone antes de WhatsApp, com pergunta de contexto antes de preço",
         "Rafael Lima", "100% dos leads em até 5 min", "Diária",
         "Relatório de tempo de primeira resposta, revisado toda segunda"],
    ]
    processos = dados.get("processos") or []
    r = 9
    if processos:
        for p in processos:
            linha = [p.get("fonte", ""), p.get("atividade", ""), p.get("padrao", ""),
                     p.get("responsavel", ""), p.get("meta_atividade", ""),
                     p.get("frequencia", ""), p.get("verificacao", "")]
            for i, v in enumerate(linha):
                cc = cel(ws, f"{get_column_letter(1+i)}{r}", v, size=9, wrap=True, border=True, valign="top")
                input_style(cc)
                cc.font = Font(name=FONTE, bold=True, size=9, color=AZUL_INPUT)
            ws.row_dimensions[r].height = 34
            r += 1
    else:
        for ex in exemplos:
            for i, v in enumerate(ex):
                cel(ws, f"{get_column_letter(1+i)}{r}", v, size=9, wrap=True, border=True,
                    italic=True, color=CINZA_ESC, valign="top")
            ws.row_dimensions[r].height = 34
            r += 1

    for _ in range(10):
        for i in range(7):
            cc = cel(ws, f"{get_column_letter(1+i)}{r}", None, size=9, wrap=True, border=True, valign="top")
            input_style(cc)
        ws.row_dimensions[r].height = 28
        r += 1

    r += 1
    faixa(ws, r, "OS 3 RITUAIS QUE SUSTENTAM A OPERAÇÃO — marque na agenda, mesmo que seja só você", "G", cor=LARANJA)
    r += 1
    cabecalho_tabela(ws, r, ["Ritual", "Frequência", "O que olhar", "Formato", "Quem conduz", "Dia e hora na agenda"])
    cel(ws, f"G{r}", "", fill=PRETO, border=True)
    r += 1
    rituais = [
        ["Matinal", "Diária · 15 min", "Prioridades do dia, travamentos, próximos passos",
         "Em pé, sem pauta longa", "", ""],
        ["Comitê de receita", "Semanal · 30 min",
         "Meta vs. realizado por fonte (aba 4), indicadores de direção, ações",
         "Tabela na tela + UMA decisão concreta", "", ""],
        ["Planejamento", "Mensal · 1 a 2h",
         "Ajustes no plano, capacidade do time, novas iniciativas",
         "Revisão do sistema completo (abas 1 a 3)", "", ""],
    ]
    for rit in rituais:
        for i, v in enumerate(rit):
            cc = cel(ws, f"{get_column_letter(1+i)}{r}", v, size=9, wrap=True, border=True, valign="top")
            if i >= 4:
                input_style(cc)
        ws.row_dimensions[r].height = 30
        r += 1

    r += 1
    faixa(ws, r, "SE VOCÊ FIZER SÓ ISSO, JÁ PAGOU O MÊS", "G")
    r += 1
    paragrafo(ws, r, "Escreva, para as suas 3 principais fontes de receita, o que precisa ser feito, quantas vezes, como e por quem. "
                     "Depois marque o comitê de receita semanal na agenda — mesmo que seja só você — e analise que atividades você "
                     "pode fazer mais e melhor para vender mais.", "G", bold=True, fill=OFFWHITE, chars=140)
    r += 2
    faixa(ws, r, "AS 7 PERGUNTAS PARA INVESTIGAR A SUA VENDA", "G", cor=LARANJA)
    r += 1
    perguntas = [
        "Quais são as fontes de receita do meu negócio hoje? Eu tenho nome para cada uma delas?",
        "Qual é a receita por entrada de cada fonte? Eu sei qual vale mais o meu tempo?",
        "Para as atividades mais críticas de vender: existe padrão, responsável e forma de verificar?",
        "Eu tenho rituais de vendas na agenda — matinal, comitê de receita semanal e planejamento mensal?",
        "Qual foi a última vez que eu mesmo atendi ou ajudei a fechar um cliente?",
        "Quantas vezes eu ofereci o meu produto ou serviço hoje?",
        "Como eu ofereço meu produto ou serviço para 10, 20 ou 30 pessoas a mais essa semana?",
    ]
    for i, p in enumerate(perguntas, 1):
        paragrafo(ws, r, f"{i}.  {p}", "G", chars=150)
        r += 1
    return ws


# --------------------------------------------------------------------------
# Abas de exemplo
# --------------------------------------------------------------------------
def aba_exemplo(wb, cfg):
    ws = wb.create_sheet(cfg["aba"])
    ws.sheet_view.showGridLines = False
    larguras(ws, {"A": 32, "B": 12, "C": 12, "D": 11, "E": 15, "F": 13,
                  "G": 13, "H": 13, "I": 13, "J": 15, "K": 12})

    titulo_aba(ws, cfg["titulo"], cfg["subtitulo"], "K")

    r = 4
    faixa(ws, r, "O NEGÓCIO", "K", cor=LARANJA)
    r += 1
    for txt in cfg["contexto"]:
        paragrafo(ws, r, txt, "K", chars=150)
        r += 1
    r += 1

    faixa(ws, r, "O PLANO DE RECEITA — NÍVEIS 1, 2 E 3", "K")
    r += 1
    cabecalho_tabela(ws, r, ["Fonte de receita", cfg["nome_entrada"], cfg["nome_proposta"], "Vendas",
                             "Receita (R$)", "Ticket médio", f"Conv. entrada→{cfg['nome_proposta'].lower()}",
                             "Conv. →venda", "Conv. entrada→venda", "Receita por entrada", "% da receita"])
    r += 1
    ini = r
    for f in cfg["fontes"]:
        cel(ws, f"A{r}", f[0], bold=True, size=9, wrap=True, border=True)
        for col, val, fmt in (("B", f[1], FMT_INT), ("C", f[2], FMT_INT),
                              ("D", f[3], FMT_INT), ("E", f[4], FMT_RS)):
            cel(ws, f"{col}{r}", val, fmt=fmt, align="center", border=True, size=9)
        cel(ws, f"F{r}", f"=IFERROR(E{r}/D{r},\"\")", fmt=FMT_RS, align="center", border=True, size=9)
        cel(ws, f"G{r}", f"=IFERROR(C{r}/B{r},\"\")", fmt=FMT_PCT, align="center", border=True, size=9)
        cel(ws, f"H{r}", f"=IFERROR(D{r}/C{r},\"\")", fmt=FMT_PCT, align="center", border=True, size=9)
        cel(ws, f"I{r}", f"=IFERROR(D{r}/B{r},\"\")", fmt=FMT_PCT, align="center", border=True, size=9)
        cel(ws, f"J{r}", f"=IFERROR(E{r}/B{r},\"\")", fmt=FMT_RS2, align="center", border=True,
            size=10, bold=True, color=LARANJA)
        cel(ws, f"K{r}", f"=IFERROR(E{r}/$E${ini+len(cfg['fontes'])},\"\")", fmt=FMT_PCT,
            align="center", border=True, size=9)
        r += 1
    fim = r - 1
    cel(ws, f"A{r}", "TOTAL", bold=True, color="FFFFFF", fill=PRETO, border=True)
    for col, fmt in (("B", FMT_INT), ("C", FMT_INT), ("D", FMT_INT), ("E", FMT_RS)):
        cel(ws, f"{col}{r}", f"=SUM({col}{ini}:{col}{fim})", bold=True, color="FFFFFF",
            fill=PRETO, align="center", border=True, fmt=fmt)
    cel(ws, f"F{r}", f"=IFERROR(E{r}/D{r},\"\")", bold=True, color="FFFFFF", fill=PRETO, align="center", border=True, fmt=FMT_RS)
    cel(ws, f"G{r}", f"=IFERROR(C{r}/B{r},\"\")", bold=True, color="FFFFFF", fill=PRETO, align="center", border=True, fmt=FMT_PCT)
    cel(ws, f"H{r}", f"=IFERROR(D{r}/C{r},\"\")", bold=True, color="FFFFFF", fill=PRETO, align="center", border=True, fmt=FMT_PCT)
    cel(ws, f"I{r}", f"=IFERROR(D{r}/B{r},\"\")", bold=True, color="FFFFFF", fill=PRETO, align="center", border=True, fmt=FMT_PCT)
    cel(ws, f"J{r}", f"=IFERROR(E{r}/B{r},\"\")", bold=True, color="FFFFFF", fill=PRETO, align="center", border=True, fmt=FMT_RS2)
    cel(ws, f"K{r}", f"=SUM(K{ini}:K{fim})", bold=True, color="FFFFFF", fill=PRETO, align="center", border=True, fmt=FMT_PCT)
    r += 2

    faixa(ws, r, "COMO LER ESSE PLANO — o que esses números estão dizendo", "K", cor=LARANJA)
    r += 1
    for i, (tit, txt) in enumerate(cfg["leitura"], 1):
        cel(ws, f"A{r}", f"{i}. {tit}", bold=True, size=10, color=LARANJA, wrap=True, valign="top", border=True)
        ws.merge_cells(f"B{r}:K{r}")
        cel(ws, f"B{r}", txt, size=9, wrap=True, valign="top", border=True)
        ws.row_dimensions[r].height = max(28, math.ceil(len(txt) / 115) * 13)
        r += 1
    r += 1

    faixa(ws, r, "A DECISÃO QUE ESSE PLANO GERA", "K")
    r += 1
    for txt in cfg["decisao"]:
        paragrafo(ws, r, "→  " + txt, "K", chars=145)
        r += 1
    r += 1

    faixa(ws, r, "SIMULAÇÃO — o que acontece se o dono puxar as alavancas certas", "K", cor=LARANJA)
    r += 1
    cabecalho_tabela(ws, r, ["Movimento", "Alavanca", "Hoje", "Depois do movimento",
                             "Impacto na receita (R$/mês)", "Custo do movimento e observação"])
    for _c in ("G", "H", "I", "J", "K"):
        cel(ws, f"{_c}{r}", "", fill=PRETO, border=True)
    ws.merge_cells(f"F{r}:K{r}")
    r += 1
    sim_ini = r
    for mov in cfg["simulacao"]:
        cel(ws, f"A{r}", mov[0], size=9, bold=True, wrap=True, border=True, valign="top")
        cel(ws, f"B{r}", mov[1], size=9, wrap=True, border=True, align="center", valign="top")
        cel(ws, f"C{r}", mov[2], size=9, wrap=True, border=True, align="center", valign="top")
        cel(ws, f"D{r}", mov[3], size=9, wrap=True, border=True, align="center", valign="top")
        cel(ws, f"E{r}", mov[4], size=10, bold=True, color=LARANJA, fmt=FMT_RS, border=True, align="center")
        ws.merge_cells(f"F{r}:K{r}")
        cel(ws, f"F{r}", mov[5], size=9, wrap=True, border=True, italic=True, valign="top")
        ws.row_dimensions[r].height = 32
        r += 1
    cel(ws, f"A{r}", "GANHO TOTAL DO MÊS", bold=True, color="FFFFFF", fill=PRETO, border=True)
    for col in ("B", "C", "D"):
        cel(ws, f"{col}{r}", "", fill=PRETO, border=True)
    cel(ws, f"E{r}", f"=SUM(E{sim_ini}:E{r-1})", bold=True, color="FFFFFF", fill=PRETO,
        align="center", border=True, fmt=FMT_RS)
    ws.merge_cells(f"F{r}:K{r}")
    cel(ws, f"F{r}", cfg["fecho_simulacao"], bold=True, color="FFFFFF", fill=PRETO, size=9, wrap=True)
    r += 2

    faixa(ws, r, "A LIÇÃO DESSE EXEMPLO", "K")
    r += 1
    paragrafo(ws, r, cfg["licao"], "K", bold=True, fill=OFFWHITE, chars=130)
    return ws


EXEMPLO_B2B = {
    "aba": "Exemplo Serviço B2B",
    "titulo": "EXEMPLO 1 · SERVIÇO B2B — ATIVA CONSULTORIA",
    "subtitulo": "Consultoria de gestão financeira para PMEs · projetos de R$ 12 mil a R$ 25 mil · time de 9 pessoas",
    "nome_entrada": "Entradas (leads)",
    "nome_proposta": "Propostas",
    "contexto": [
        "A Ativa vende projetos de estruturação financeira para empresas de R$ 5 a R$ 50 milhões de faturamento. Fatura cerca de R$ 357 mil por mês e cresce, mas o sócio não sabe explicar de onde vem o crescimento — nos meses bons ele diz «foi um mês bom».",
        "O time comercial tem 1 sócio vendendo, 1 executivo de contas e 1 SDR fazendo prospecção ativa. A sensação do sócio, antes do plano: «prospecção ativa é o motor da empresa, é onde a gente mais trabalha». Os números abaixo mostram outra coisa.",
        "Período do plano: um mês fechado. Entrada = qualquer contato inicial (formulário do site, resposta de e-mail frio, indicação recebida, contato em evento). Proposta = documento comercial enviado com escopo e preço.",
    ],
    "fontes": [
        ["Inbound (site, Google, conteúdo)", 110, 22, 5, 90000],
        ["Indicação de clientes", 8, 6, 4, 100000],
        ["Prospecção ativa (SDR)", 140, 14, 2, 44000],
        ["Carteira ativa (expansão de escopo)", 20, 10, 5, 60000],
        ["Parcerias (contadores e bancos)", 12, 7, 2, 44000],
        ["Eventos e palestras", 30, 6, 1, 19000],
    ],
    "leitura": [
        ("Nível 1 · A maior fonte não é a que parece",
         "Indicação gera R$ 100 mil por mês — a maior fatia da receita (28%) — com apenas 8 entradas. Prospecção ativa, onde o time gasta a maior parte do dia, gera R$ 44 mil (12%). O sócio achava que o motor era a prospecção porque é onde há mais movimento, não mais dinheiro. Movimento não é receita."),
        ("Nível 2 · O gargalo do inbound está no meio do funil",
         "O inbound recebe 110 entradas e transforma só 22 em proposta (20%). São 88 pessoas por mês que levantaram a mão e nunca receberam um preço. Não é problema de geração de lead — é problema de atendimento. Comprar mais lead antes de arrumar isso é jogar dinheiro em um funil furado."),
        ("Nível 2 · A prospecção não está errada, está mal calibrada",
         "140 entradas viram 14 propostas e 2 vendas: 1,4% de conversão entrada→venda, a pior do plano. O problema não é o SDR trabalhar pouco — é a lista e a abordagem. Antes de contratar o segundo SDR, o dono precisa arrumar quem esse SDR aborda e com qual frase."),
        ("Nível 3 · Uma indicação vale 40 prospecções",
         "Cada entrada de indicação gera R$ 12.500 de receita. Cada entrada de prospecção ativa gera R$ 314. É 40 vezes mais. Se o sócio tem 1 hora livre para vender, ela rende 40x mais pedindo indicação do que fazendo ligação fria. Esse é o número que muda a agenda do dono."),
        ("Nível 3 · Inbound é volume, indicação é valor — e cada um pede uma ação diferente",
         "Inbound tem R$ 818 por entrada e 110 entradas. Melhorar R$ 100 na receita por entrada do inbound vale R$ 11 mil por mês, sem nenhuma entrada nova. Indicação tem poucas entradas e valor altíssimo: aqui o jogo é gerar mais entradas, não otimizar. Canal de volume se otimiza; canal de valor se multiplica."),
        ("O ponto cego",
         "Indicação é a maior fonte de receita da empresa e é a única sem nenhum processo: acontece quando o cliente lembra. Uma fonte de R$ 100 mil por mês rodando por sorte é o maior risco do negócio — e, ao mesmo tempo, a maior oportunidade."),
    ],
    "decisao": [
        "Transformar indicação em processo, não em sorte: pedir 2 nomes no call de encerramento de todo projeto, com responsável, meta (100% das entregas) e verificação no CRM. Meta: sair de 8 para 16 entradas por mês.",
        "Arrumar o meio do funil do inbound antes de comprar mais tráfego: responder todo lead em até 5 minutos e subir a conversão entrada→proposta de 20% para 35%.",
        "Congelar a expansão da prospecção ativa por 60 dias. Não cortar — recalibrar lista e abordagem, e só voltar a escalar quando a conversão entrada→venda passar de 1,4% para 3%.",
        "Manter tudo que já funciona. Crescer é empilhar fontes que funcionam, não trocar a aposta: a carteira ativa continua com a mesma cadência de contato.",
    ],
    "simulacao": [
        ["Dobrar as indicações", "Mais entradas", "8 entradas/mês", "16 entradas/mês", 100000,
         "Custo: R$ 0. Só um processo no encerramento de projeto. É a receita mais barata do plano inteiro."],
        ["Arrumar o funil do inbound", "Melhor conversão", "20% entrada→proposta", "35% entrada→proposta", 67500,
         "Custo: uma regra de resposta em 5 minutos. As 110 entradas já foram pagas — só estavam sendo desperdiçadas."],
        ["Recalibrar a prospecção ativa", "Melhor conversão", "1,4% entrada→venda", "3,0% entrada→venda", 48400,
         "Custo: tempo de revisão de lista e script. Sem contratar o segundo SDR."],
    ],
    "fecho_simulacao": "De R$ 357 mil para R$ 573 mil por mês sem contratar ninguém e sem aumentar o investimento em mídia.",
    "licao": "O plano de receita não inventou nenhum canal novo. Ele só deu nome ao que já existia — e mostrou que a empresa estava colocando o maior esforço (prospecção) na entrada que vale 40 vezes menos, enquanto a fonte mais rentável (indicação) rodava sem processo, sem dono e sem meta. É isso que fazer gestão de receita significa: saber de onde o dinheiro vem, olhar cada fonte separadamente e montar um plano específico para cada uma.",
}

EXEMPLO_VAREJO = {
    "aba": "Exemplo Loja Física",
    "titulo": "EXEMPLO 2 · VAREJO — LOJA ORIGEM",
    "subtitulo": "Loja de moda masculina de rua · 1 unidade · 4 vendedores · faturamento de R$ 142 mil/mês",
    "nome_entrada": "Entradas (fluxo)",
    "nome_proposta": "Atendimentos",
    "contexto": [
        "A Origem é uma loja de rua com boa localização, 4 vendedores e 9 anos de estrada. O dono acompanha faturamento e ticket médio, mas nunca separou a receita por origem — para ele, «venda é venda». Quando o mês cai, a explicação é sempre a mesma: «o movimento da rua caiu».",
        "Ao separar as fontes, ele descobriu que só metade da receita depende do movimento da rua — e que a metade que ele controla é justamente a que ninguém cuida.",
        "Período do plano: um mês fechado. Entrada = pessoa que entra na loja, conversa iniciada no WhatsApp, clique no anúncio, visita ao e-commerce ou contato de empresa. Atendimento = cliente que experimentou peça, pediu orçamento ou colocou item no carrinho. No varejo, «proposta» vira «atendimento» — o nome muda, a lógica é a mesma.",
    ],
    "fontes": [
        ["Fluxo de rua (quem passa e entra)", 1400, 560, 224, 71680],
        ["Base de clientes no WhatsApp", 260, 130, 78, 35100],
        ["Instagram e anúncio local", 180, 72, 25, 8750],
        ["Indicação de cliente", 40, 28, 18, 7560],
        ["Empresas e uniformes (B2B)", 12, 6, 3, 12600],
        ["E-commerce e marketplace", 900, 90, 22, 6380],
    ],
    "leitura": [
        ("Nível 1 · «Movimento caiu» explica metade do negócio, não o negócio inteiro",
         "O fluxo de rua responde por R$ 71,7 mil (50% da receita). A outra metade — WhatsApp da base, indicação, B2B, e-commerce — não depende de quem passa na calçada. Enquanto tudo era «venda é venda», o dono culpava a rua por um resultado que ele controlava em 50%."),
        ("Nível 2 · A conversão da loja é o ativo mais valioso e o menos gerenciado",
         "De 1.400 pessoas que entram, 560 são atendidas de verdade (40%) e 224 compram (16% do fluxo total). Ou seja: 840 pessoas por mês entram na loja e saem sem serem atendidas. Cada ponto percentual de conversão do fluxo vale cerca de R$ 4,5 mil por mês — sem gastar R$ 1 a mais de anúncio."),
        ("Nível 3 · Uma entrada B2B vale 20 fluxos de rua",
         "Cada contato de empresa gera R$ 1.050 de receita. Cada pessoa que entra na loja gera R$ 51. É 20 vezes mais. E o canal B2B teve 12 entradas no mês — todas passivas, ninguém foi atrás de nenhuma. É o canal mais rentável da loja rodando no automático."),
        ("Nível 3 · A base no WhatsApp é a alavanca mais barata que existe",
         "R$ 135 por entrada, ticket de R$ 450 (40% acima do fluxo de rua) e conversão de 30% — o dobro da loja. Faz sentido: é gente que já comprou e já confia. E essa lista já está no celular dos vendedores, sem custo de aquisição nenhum."),
        ("Nível 3 · O e-commerce está consumindo esforço e devolvendo centavos",
         "900 visitas geram R$ 6,4 mil e R$ 7,09 por entrada — o pior do plano, com 2,4% de conversão. Isso não significa fechar. Significa parar de tratá-lo como aposta de crescimento enquanto a conversão não sair do lugar, e não colocar mais verba de mídia ali antes de arrumar o funil."),
        ("O ponto cego",
         "O dono estava prestes a contratar uma agência para «trazer mais gente para a loja». O plano mostrou que trazer mais gente é a alavanca mais cara — e que ele tinha 840 pessoas por mês entrando sem atendimento, uma lista de clientes parada e um canal B2B 20 vezes mais rentável sem dono."),
    ],
    "decisao": [
        "Definir processo de atendimento com padrão verificável: abordar em até 2 minutos, a 3 metros da porta, com pergunta de contexto («é a primeira vez aqui?») e oferta do produto-âncora em todo atendimento. Meta: subir a conversão do fluxo de 16% para 18%.",
        "Dar dono ao canal B2B: um vendedor responsável por 20 contatos ativos por mês com empresas do bairro (escritórios, clínicas, concessionárias) — o canal mais rentável não pode continuar passivo.",
        "Criar cadência de WhatsApp ativo: cada vendedor faz 15 conversas por dia com a base, com motivo real (peça nova do tamanho dele, aniversário, item que combina com a última compra). Meta: de 260 para 400 entradas por mês.",
        "Segurar a verba nova de anúncio até a conversão da loja subir. Não faz sentido pagar para trazer gente que ninguém vai atender.",
    ],
    "simulacao": [
        ["Processo de atendimento na loja", "Melhor conversão", "16% do fluxo compra", "18% do fluxo compra", 8960,
         "Custo: treino de 1 hora e cobrança diária na matinal. As 1.400 entradas já estavam lá."],
        ["Dar dono ao canal B2B", "Mais entradas", "12 contatos/mês", "24 contatos/mês", 12600,
         "Custo: R$ 0 de mídia. Um vendedor com meta de 20 contatos ativos por mês."],
        ["Cadência de WhatsApp com a base", "Mais entradas", "260 conversas/mês", "400 conversas/mês", 18900,
         "Custo: R$ 0. A lista já existe. É a receita que estava dormindo no celular do time."],
        ["Peça adicional no fechamento", "Ticket maior", "Ticket de R$ 320", "Ticket de R$ 350", 6720,
         "Custo: uma pergunta padrão no caixa. Vale para as 224 vendas do fluxo de rua."],
    ],
    "fecho_simulacao": "De R$ 142 mil para R$ 189 mil por mês — 33% a mais — sem um real a mais de anúncio e sem contratar ninguém.",
    "licao": "No varejo, a tentação é sempre resolver venda com movimento: mais anúncio, mais gente na porta, mais promoção. O plano de receita mostra que as alavancas mais baratas quase sempre estão dentro de casa — na conversão de quem já entrou, na base que já comprou e no canal rentável que ninguém está olhando. Não existe mês ruim quando você sabe de onde a venda vem: existe alavanca que ninguém puxou.",
}


# --------------------------------------------------------------------------
# ABA — Glossário
# --------------------------------------------------------------------------
def aba_glossario(wb):
    ws = wb.create_sheet("Glossário")
    ws.sheet_view.showGridLines = False
    larguras(ws, {"A": 28, "B": 60, "C": 60})

    titulo_aba(ws, "GLOSSÁRIO E MÉTODO", "Os termos do plano de receita, em português de dono", "C")

    r = 4
    faixa(ws, r, "OS TERMOS", "C", cor=LARANJA)
    r += 1
    cabecalho_tabela(ws, r, ["Termo", "O que é", "Exemplo no seu negócio"])
    r += 1
    termos = [
        ["Fonte de receita", "A origem da venda. Cada fonte é, na prática, um dinheiro diferente: exige esforço diferente, converte diferente e vale diferente.",
         "Carteira ativa, indicação, prospecção ativa, inbound, eventos, parcerias, canal indireto, fluxo de loja, WhatsApp, marketplace."],
        ["Entrada", "Todo contato inicial, seja qual for o formato. É o termo genérico que serve para qualquer negócio.",
         "Pessoa que entra na loja · lead que preenche formulário · orçamento pedido · visita agendada · DM respondida · ligação atendida."],
        ["Proposta", "O meio do caminho entre a entrada e a venda. Se o seu negócio não tem essa etapa, adapte o nome — o que importa é medir o meio do funil.",
         "Proposta enviada · orçamento passado · cliente que experimentou a peça · degustação · aula experimental · carrinho montado."],
        ["Venda", "O negócio fechado, atribuído à fonte que o originou.", "Contrato assinado · cupom emitido · pedido faturado."],
        ["Conversão", "Quanto de uma etapa passa para a próxima. É o indicador de direção que mostra onde a venda trava.",
         "100 entradas → 20 propostas = 20% de conversão entrada→proposta."],
        ["Ticket médio", "Receita dividida por vendas. Quanto vale, em média, cada negócio fechado naquela fonte.",
         "R$ 90.000 de receita ÷ 5 vendas = R$ 18.000 de ticket médio."],
        ["Receita por entrada", "Receita da fonte dividida pelas entradas da fonte. É o indicador do nível 3 — o que quase ninguém calcula e o que muda a decisão.",
         "R$ 100.000 ÷ 8 indicações = R$ 12.500 por entrada. Contra R$ 314 da prospecção ativa."],
        ["Indicador de direção", "O número que você monitora antes do resultado aparecer: conversão, receita por entrada, cobertura, ciclo.",
         "Se as atividades acontecem e os indicadores de direção não se movem, o processo precisa ser revisto."],
        ["Atividade", "Tudo o que está no seu controle fazer. Fazer alguém entrar na loja não está; atendê-la depois que entrou, sim.",
         "10 ligações por dia · responder em 5 minutos · pedir 2 indicações por entrega."],
        ["Processo", "Atividade com padrão, responsável e forma de verificar. Sem os 3, é intenção — e intenção não se ensina, não se mede e não se cobra.",
         "«Atender bem» é intenção. «Receber a 3 metros da porta em até 2 minutos e perguntar o que trouxe a pessoa» é processo."],
    ]
    for t in termos:
        cel(ws, f"A{r}", t[0], bold=True, size=10, color=LARANJA, wrap=True, valign="top", border=True)
        cel(ws, f"B{r}", t[1], size=9, wrap=True, valign="top", border=True)
        cel(ws, f"C{r}", t[2], size=9, wrap=True, valign="top", italic=True, border=True)
        ws.row_dimensions[r].height = max(32, math.ceil(max(len(t[1]), len(t[2])) / 75) * 13)
        r += 1

    r += 1
    faixa(ws, r, "OS 3 ESTÁGIOS DE VENDAS — em qual você está?", "C")
    r += 1
    cabecalho_tabela(ws, r, ["Estágio", "A pergunta que você faz", "O que define"])
    r += 1
    estagios = [
        ["01 · Reativo", "Quanto você está vendendo?",
         "Sem plano, sem processos, sem cadência. O resultado depende do esforço individual e da rede de contatos do dono. Funciona, mas é frágil."],
        ["02 · Estruturado", "Quanto você quer vender?",
         "Plano por fonte de receita, processos-chave definidos, cadência semanal de reunião e diagnóstico individual do time."],
        ["03 · Escalável", "Quais as principais ações para replicar as vendas de diferentes origens?",
         "Análise por canal, remuneração atrelada ao plano, playbooks por perfil e sistema integrado. A empresa opera vários silos de receita ao mesmo tempo."],
    ]
    for e in estagios:
        cel(ws, f"A{r}", e[0], bold=True, size=10, color=LARANJA, wrap=True, valign="top", border=True)
        cel(ws, f"B{r}", e[1], size=9, wrap=True, valign="top", border=True)
        cel(ws, f"C{r}", e[2], size=9, wrap=True, valign="top", border=True)
        ws.row_dimensions[r].height = 40
        r += 1
    r += 1
    paragrafo(ws, r, "Um mesmo negócio pode estar em estágios diferentes em cada frente: a loja física madura e o digital ainda engatinhando. "
                     "Reconhecer em que fase você está, em cada frente, é o primeiro passo.", "C", italic=True, chars=110)
    return ws


# --------------------------------------------------------------------------
def gerar(dados, saida):
    wb = Workbook()
    wb.remove(wb.active)

    aba_comece_aqui(wb, dados)
    aba_plano(wb, dados)
    aba_simulador(wb, dados)
    aba_projecao(wb, dados)
    aba_acompanhamento(wb, dados)
    aba_processos(wb, dados)
    aba_exemplo(wb, EXEMPLO_B2B)
    aba_exemplo(wb, EXEMPLO_VAREJO)
    aba_glossario(wb)

    for ws in wb.worksheets:
        ws.freeze_panes = "A3"
    # Marca o arquivo para recalcular todas as fórmulas ao abrir.
    wb.calculation.fullCalcOnLoad = True
    wb.save(saida)
    return saida


def main():
    args = sys.argv[1:]
    if not args:
        print(__doc__)
        sys.exit(1)
    if args[0] == "--modelo":
        dados = {}
        saida = args[1] if len(args) > 1 else "Plano_de_Receita_MODELO.xlsx"
    else:
        with open(args[0], encoding="utf-8") as fh:
            dados = json.load(fh)
        saida = args[1] if len(args) > 1 else "Plano_de_Receita.xlsx"
    gerar(dados, saida)
    print(f"OK: {saida}")
    if dados.get("fontes"):
        try:
            import os
            sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
            from calcular_plano import calcular, markdown
            print()
            print("Resumo calculado (mesmas fórmulas da aba 1. Plano de Receita):")
            print(markdown(calcular(dados), dados))
        except Exception as exc:  # o resumo é um extra; a planilha já foi salva
            print(f"(Resumo não calculado: {exc})")


if __name__ == "__main__":
    main()
