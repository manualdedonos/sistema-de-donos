#!/usr/bin/env python3
"""
Gera a planilha "Calendário Comercial" (.xlsx) a partir de um JSON.

Uso:
    python3 gerar_calendario_xlsx.py dados.json "Calendario_Comercial_Negocio.xlsx"

Requer: Python 3 e openpyxl.

Sempre gera as 4 abas: Calendário, Datas Próprias, Plano de Preparação e
Indicadores. Se "datas_proprias" ou "preparacao" vierem vazios, a aba é
criada mesmo assim: Datas Próprias com uma linha "a preencher" e Plano de
Preparação com uma linha por ativação do calendário (campanha e início da
preparação preenchidos a partir do calendário, demais etapas "a definir").

----------------------------------------------------------------------
FORMATO DO JSON DE ENTRADA
----------------------------------------------------------------------
Todos os campos de texto são livres. Só "calendario" é obrigatório; as
demais seções são opcionais mas recomendadas (sem elas, as abas saem com
linhas "a preencher" / "a definir").

{
  "negocio": "Farmácia São João",            # nome do negócio (cabeçalho)
  "diagnostico": "Farmácia de bairro em ...", # resumo do diagnóstico (1-2 frases)
  "objetivo": "Aumentar frequência de compra da base",
  "periodo": "12 meses",

  "calendario": [                             # OBRIGATÓRIO — uma linha por ativação
    {
      "mes": "Janeiro",
      "semana": "Semana 2",
      "camada": "própria",                    # nacional | regional | setor | própria
      "data": "Virada de ano - Recomeço",
      "objetivo": "Reativar a base",
      "conceito": "Kit Início de Ano",
      "oferta": "Combo imunidade com 15% para a base",
      "canal": "WhatsApp",
      "preparar": "Dezembro (semana 3)",
      "indicador": "Recompra da base vs. jan normal"
    }
  ],

  "datas_proprias": [                         # opcional — pelo menos 3 recomendado
    {
      "nome": "Clube da Farmácia - Dia do Cliente",
      "periodo": "Todo dia 15",
      "porque": "Cria motivo recorrente de volta; ataca a frequência",
      "vale": "Baixa recorrência entre necessidades",
      "oferta": "Condição especial para a base, avisada no WhatsApp"
    }
  ],

  "preparacao": [                             # opcional — trabalho de trás pra frente
    {
      "campanha": "Outubro Rosa",
      "definir_oferta": "Set, semana 1",
      "estoque": "Set, semana 3",
      "alinhar_time": "Set, semana 4",
      "iniciar_comunicacao": "Out, semana 1",
      "medir": "Out, semana 5"
    }
  ],

  "indicadores": [                            # opcional — como medir sucesso
    {
      "indicador": "Frequência de compra",
      "o_que_mede": "Quantas vezes o cliente volta no ano",
      "meta": "Subir de 4 para 5 compras/ano"
    }
  ]
}
----------------------------------------------------------------------
"""

import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Paleta
AZUL = "1F3A5F"
AZUL_CLARO = "DCE6F1"
CINZA = "F2F2F2"
BORDA = Side(style="thin", color="BFBFBF")
ALL_BORDERS = Border(left=BORDA, right=BORDA, top=BORDA, bottom=BORDA)

CAMADA_CORES = {
    "nacional": "FFE2CC",   # laranja claro
    "regional": "E2EFDA",   # verde claro
    "setor": "FFF2CC",      # amarelo claro
    "própria": "DCE6F1",    # azul claro
    "propria": "DCE6F1",
}


def _style_header(cell):
    cell.font = Font(bold=True, color="FFFFFF", size=11)
    cell.fill = PatternFill("solid", fgColor=AZUL)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = ALL_BORDERS


def _title_block(ws, titulo, subtitulo=None, ncols=10):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=titulo)
    c.font = Font(bold=True, size=16, color=AZUL)
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26
    start = 2
    if subtitulo:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
        c = ws.cell(row=2, column=1, value=subtitulo)
        c.font = Font(italic=True, size=11, color="595959")
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[2].height = 20
        start = 3
    return start


def _write_table(ws, header_row, headers, rows, widths, camada_col=None):
    for j, h in enumerate(headers, start=1):
        _style_header(ws.cell(row=header_row, column=j, value=h))
    ws.row_dimensions[header_row].height = 30
    for i, row in enumerate(rows):
        r = header_row + 1 + i
        for j, key in enumerate(headers_keys(headers), start=1):
            val = row.get(key, "")
            cell = ws.cell(row=r, column=j, value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = ALL_BORDERS
            if i % 2 == 1:
                cell.fill = PatternFill("solid", fgColor=CINZA)
        # colorir por camada
        if camada_col is not None:
            camada = str(row.get("camada", "")).strip().lower()
            cor = CAMADA_CORES.get(camada)
            if cor:
                cl = ws.cell(row=r, column=camada_col)
                cl.fill = PatternFill("solid", fgColor=cor)
                cl.font = Font(bold=True)
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)


# Mapeamento header->chave por aba (preenchido nas funções)
_KEYS = []


def headers_keys(headers):
    return _KEYS


def aba_calendario(wb, dados):
    global _KEYS
    ws = wb.active
    ws.title = "Calendário"
    start = _title_block(
        ws,
        f"Calendário Comercial — {dados.get('negocio', '')}".strip(" —"),
        dados.get("diagnostico", ""),
        ncols=10,
    )
    headers = [
        "Mês", "Semana aprox.", "Camada", "Data / oportunidade",
        "Objetivo da ativação", "Conceito da campanha", "Oferta / mecânica",
        "Canal principal", "Quando preparar", "Indicador de sucesso",
    ]
    _KEYS = ["mes", "semana", "camada", "data", "objetivo", "conceito",
             "oferta", "canal", "preparar", "indicador"]
    widths = [12, 13, 11, 26, 24, 22, 28, 14, 16, 24]
    _write_table(ws, start, headers, dados.get("calendario", []), widths, camada_col=3)


def aba_datas_proprias(wb, dados):
    global _KEYS
    linhas = dados.get("datas_proprias") or [
        {"nome": "A preencher", "periodo": "", "porque": "Proponha pelo menos 3 datas próprias, de preferência sobre os vales.",
         "vale": "", "oferta": ""},
    ]
    ws = wb.create_sheet("Datas Próprias")
    start = _title_block(ws, "Datas Próprias — as que você cria do zero",
                         "Onde não há concorrente: você define a regra, a oferta e a história.", ncols=5)
    headers = ["Nome da data", "Melhor período", "Por que faz sentido",
               "Vale que preenche", "Oferta / experiência"]
    _KEYS = ["nome", "periodo", "porque", "vale", "oferta"]
    widths = [26, 18, 34, 26, 34]
    _write_table(ws, start, headers, linhas, widths)


def aba_preparacao(wb, dados):
    global _KEYS
    linhas = dados.get("preparacao")
    if not linhas:
        # Sem bloco de preparação: uma linha por ativação, etapas a definir.
        linhas = []
        for item in dados.get("calendario", []):
            linhas.append({
                "campanha": item.get("data") or item.get("conceito") or "A preencher",
                # "preparar" = quando começar a preparar: é a primeira etapa.
                "definir_oferta": item.get("preparar") or "a definir",
                "estoque": "a definir",
                "alinhar_time": "a definir",
                "iniciar_comunicacao": "a definir",
                "medir": "a definir",
            })
        if not linhas:
            linhas = [{"campanha": "A preencher"}]
    ws = wb.create_sheet("Plano de Preparação")
    start = _title_block(ws, "Plano de Preparação — de trás para frente",
                         "A venda boa nasce semanas antes da data.", ncols=6)
    headers = ["Campanha", "Definir oferta", "Estoque/agenda pronto",
               "Alinhar time", "Iniciar comunicação", "Medir resultado"]
    _KEYS = ["campanha", "definir_oferta", "estoque", "alinhar_time",
             "iniciar_comunicacao", "medir"]
    widths = [24, 18, 20, 16, 20, 18]
    _write_table(ws, start, headers, linhas, widths)


def aba_indicadores(wb, dados):
    global _KEYS
    indicadores = dados.get("indicadores") or [
        {"indicador": "Receita", "o_que_mede": "Quanto entrou no período",
         "meta": "Cresceu na semana da ativação vs. semana normal?"},
        {"indicador": "Ticket médio", "o_que_mede": "Quanto cada cliente gasta por compra",
         "meta": "Os combos subiram o ticket ou só deram desconto?"},
        {"indicador": "Frequência de compra", "o_que_mede": "Quantas vezes o cliente volta no ano",
         "meta": "Está subindo mês a mês? É o coração do calendário."},
        {"indicador": "Clientes ativos", "o_que_mede": "Quantos compraram no período",
         "meta": "A base cresce ou só gira os mesmos?"},
        {"indicador": "Recompra", "o_que_mede": "Quantos voltaram depois da 1ª compra",
         "meta": "A data própria está trazendo gente de volta?"},
    ]
    ws = wb.create_sheet("Indicadores")
    start = _title_block(ws, "Indicadores de Acompanhamento",
                         "Olhe todo mês. O calendário deixa de ser documento morto e vira sistema que aprende.", ncols=3)
    headers = ["Indicador", "O que mede", "O que olhar todo mês"]
    _KEYS = ["indicador", "o_que_mede", "meta"]
    widths = [24, 36, 50]
    _write_table(ws, start, headers, indicadores, widths)


def main():
    if len(sys.argv) < 3:
        print("Uso: python3 gerar_calendario_xlsx.py dados.json saida.xlsx")
        sys.exit(1)
    json_path, out_path = sys.argv[1], sys.argv[2]
    with open(json_path, "r", encoding="utf-8") as f:
        dados = json.load(f)
    if not dados.get("calendario"):
        print("ERRO: o JSON precisa ter ao menos a chave 'calendario' com linhas.")
        sys.exit(1)

    wb = Workbook()
    aba_calendario(wb, dados)
    aba_datas_proprias(wb, dados)
    aba_preparacao(wb, dados)
    aba_indicadores(wb, dados)
    wb.save(out_path)
    print(f"OK: planilha gerada em {out_path}")
    print(f"Abas: {wb.sheetnames}")
    print(f"Ativações no calendário: {len(dados['calendario'])}")


if __name__ == "__main__":
    main()
